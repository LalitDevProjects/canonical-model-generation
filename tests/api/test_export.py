"""Direct unit tests for api/export.py's docx/xlsx rendering, including
the declared-losses branch (empty in every tests/api/test_workshop.py
scenario, since no mapping spec is ever assembled there)."""

from __future__ import annotations

import io

import docx
import openpyxl
import pytest

from emit.workshop_pack import PackFile, WorkshopPackManifest

from api.export import render_docx, render_xlsx

pytestmark = pytest.mark.unit


def _manifest_with_losses() -> WorkshopPackManifest:
    return WorkshopPackManifest(
        domain="claims",
        version="1.0",
        files=[PackFile(path="Claim.json", sha256="a" * 64, description="Claim schema")],
        declared_losses=[{"mappingSpec": "uk-claims", "path": "ClaimHeader.LossDate", "kind": "precision", "detail": "time discarded"}],
    )


def _manifest_without_losses() -> WorkshopPackManifest:
    return WorkshopPackManifest(domain="claims", version="1.0", files=[], declared_losses=[])


class TestRenderDocx:
    def test_with_declared_losses_produces_two_tables(self) -> None:
        content = render_docx(_manifest_with_losses())
        document = docx.Document(io.BytesIO(content))
        assert len(document.tables) == 2
        loss_table = document.tables[1]
        assert loss_table.rows[1].cells[0].text == "uk-claims"

    def test_without_declared_losses_still_produces_a_document(self) -> None:
        content = render_docx(_manifest_without_losses())
        document = docx.Document(io.BytesIO(content))
        assert len(document.tables) == 1
        assert any("None." in p.text for p in document.paragraphs)


class TestRenderXlsx:
    def test_with_declared_losses_populates_the_losses_sheet(self) -> None:
        content = render_xlsx(_manifest_with_losses())
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        losses_sheet = workbook["Declared losses"]
        assert [c.value for c in losses_sheet[2]] == ["uk-claims", "ClaimHeader.LossDate", "precision", "time discarded"]

    def test_without_declared_losses_leaves_only_the_header_row(self) -> None:
        content = render_xlsx(_manifest_without_losses())
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        losses_sheet = workbook["Declared losses"]
        assert losses_sheet.max_row == 1

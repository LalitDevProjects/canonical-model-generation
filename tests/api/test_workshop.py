"""
Hermetic tests for api/workshop.py - Section 12.4's workshop and decision
API, over a real fastapi.testclient.TestClient. create_workshop() calls
the real emit.workshop_pack.assemble_pack() against a real orchestrator
run's own persisted clusters/candidates - the same real pack-assembly
path tests/emit/test_i9_acceptance.py already proves at the emit layer.
"""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from config.settings import load_settings

from api.app import create_app

pytestmark = pytest.mark.e2e


def _client(tmp_path: Path) -> TestClient:
    app = create_app(settings=load_settings(), base_path=tmp_path, require_auth=False)
    return TestClient(app)


def _create_run(client: TestClient) -> str:
    response = client.post(
        "/v1/runs",
        json={"domain": "claims", "trigger": {"kind": "manual", "requestedBy": "tester", "at": "2026-08-31T12:00:00Z"}},
    )
    assert response.status_code == 201, response.text
    run_id: str = response.json()["runId"]
    return run_id


def _create_workshop(client: TestClient, run_id: str) -> str:
    response = client.post(
        "/v1/workshops",
        json={"runId": run_id, "domain": "claims", "checkpoint": "RATIFY", "participants": [{"name": "Jane", "region": "uk", "role": "sme"}]},
    )
    assert response.status_code == 201, response.text
    workshop_id: str = response.json()["workshopId"]
    return workshop_id


def _seed_bare_manifest(tmp_path: Path, run_id: str) -> None:
    from uuid import UUID

    from generated.C11.RunManifest._1_0 import C11Runmanifest

    from pipeline.run_store import RunStore

    store = RunStore(base_path=tmp_path / "run-store")
    store.write_run_manifest(UUID(run_id), C11Runmanifest.model_validate({
        "runId": run_id, "domain": "claims",
        "trigger": {"kind": "manual", "requestedBy": "tester", "at": "2026-08-31T12:00:00Z"},
        "corpusHash": "a" * 64,
        "pins": {"prompts": {}, "models": {}, "tools": {}, "algorithms": {}},
        "parameters": {}, "budget": {"tokensTotal": 1000, "costCeilingGbp": 10.0, "perStage": {}},
        "state": "RUNNING", "createdBy": "tester",
    }))


class TestCreateWorkshop:
    def test_unknown_run_is_404(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        response = client.post(
            "/v1/workshops",
            json={"runId": "00000000-0000-0000-0000-000000000000", "domain": "claims", "checkpoint": "RATIFY", "participants": []},
        )
        assert response.status_code == 404
        assert response.json()["type"].endswith("/not-found")

    def test_run_that_exists_but_has_not_clustered_yet_is_invalid_contract(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        run_id = "11111111-1111-1111-1111-111111111111"
        _seed_bare_manifest(tmp_path, run_id)
        response = client.post(
            "/v1/workshops",
            json={"runId": run_id, "domain": "claims", "checkpoint": "RATIFY", "participants": []},
        )
        assert response.status_code == 400
        assert response.json()["type"].endswith("/invalid-contract")

    def test_returns_a_real_pack_ref_after_a_real_run(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        run_id = _create_run(client)
        response = client.post(
            "/v1/workshops",
            json={"runId": run_id, "domain": "claims", "checkpoint": "RATIFY", "participants": []},
        )
        assert response.status_code == 201
        body = response.json()
        assert body["packRef"] == f"/v1/workshops/{body['workshopId']}/pack"


class TestGetPack:
    def test_unknown_workshop_is_404(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        response = client.get("/v1/workshops/00000000-0000-0000-0000-000000000000/pack")
        assert response.status_code == 404

    def test_pack_reflects_the_runs_real_state(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        run_id = _create_run(client)
        workshop_id = _create_workshop(client, run_id)

        response = client.get(f"/v1/workshops/{workshop_id}/pack")
        assert response.status_code == 200
        body = response.json()
        # No candidates exist yet (the hermetic run never reaches S6
        # without a model provider) - an honest, empty list, not faked.
        assert body["candidates"] == []
        assert body["conflicts"] == []
        assert body["openQuestions"] == []
        assert body["coveragePreview"]["filesInPack"] > 0


class TestSubmitDecisions:
    def test_unknown_workshop_is_404(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        response = client.post(
            "/v1/workshops/00000000-0000-0000-0000-000000000000/decisions",
            json={"decisions": [{"candidateId": "canon://Claim.x", "outcome": "accept", "rationale": "r", "sme": "Jane"}]},
        )
        assert response.status_code == 404

    def test_recorded_count_matches_submitted_decisions(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        run_id = _create_run(client)
        workshop_id = _create_workshop(client, run_id)

        response = client.post(
            f"/v1/workshops/{workshop_id}/decisions",
            json={"decisions": [
                {"candidateId": "canon://Claim.a", "outcome": "accept", "rationale": "r1", "sme": "Jane"},
                {"candidateId": "canon://Claim.b", "outcome": "reject", "rationale": "r2", "sme": "Jane"},
            ]},
        )
        assert response.status_code == 201
        assert response.json()["recorded"] == 2


class TestExport:
    def test_unknown_workshop_is_404(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        response = client.get("/v1/workshops/00000000-0000-0000-0000-000000000000/export", params={"format": "docx"})
        assert response.status_code == 404

    def test_unsupported_format_is_invalid_contract(self, tmp_path: Path) -> None:
        client = _client(tmp_path)
        run_id = _create_run(client)
        workshop_id = _create_workshop(client, run_id)
        response = client.get(f"/v1/workshops/{workshop_id}/export", params={"format": "pdf"})
        assert response.status_code == 400
        assert response.json()["type"].endswith("/invalid-contract")

    def test_docx_export_is_a_real_readable_document(self, tmp_path: Path) -> None:
        import io

        client = _client(tmp_path)
        run_id = _create_run(client)
        workshop_id = _create_workshop(client, run_id)

        response = client.get(f"/v1/workshops/{workshop_id}/export", params={"format": "docx"})
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        document = docx.Document(io.BytesIO(response.content))
        assert len(document.tables) > 0

    def test_xlsx_export_is_a_real_readable_workbook(self, tmp_path: Path) -> None:
        import io

        client = _client(tmp_path)
        run_id = _create_run(client)
        workshop_id = _create_workshop(client, run_id)

        response = client.get(f"/v1/workshops/{workshop_id}/export", params={"format": "xlsx"})
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        assert "Files" in workbook.sheetnames

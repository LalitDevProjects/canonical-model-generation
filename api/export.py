"""
Renders an already-real, already-assembled workshop pack
(emit.workshop_pack.WorkshopPackManifest) into .docx / .xlsx - pure
rendering, no new business logic. Section 12.4: "GET .../export?
format=docx|xlsx -> the pack as a circulable document."
"""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook

from emit.workshop_pack import WorkshopPackManifest

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def render_docx(manifest: WorkshopPackManifest) -> bytes:
    document = Document()
    document.add_heading(f"{manifest.domain} workshop pack ({manifest.version})", level=1)

    document.add_heading("Files", level=2)
    table = document.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text = "Path", "SHA-256", "Description"
    for pack_file in manifest.files:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = pack_file.path, pack_file.sha256, pack_file.description

    document.add_heading("Declared losses", level=2)
    if manifest.declared_losses:
        loss_table = document.add_table(rows=1, cols=4)
        loss_header = loss_table.rows[0].cells
        loss_header[0].text, loss_header[1].text = "Mapping spec", "Path"
        loss_header[2].text, loss_header[3].text = "Kind", "Detail"
        for loss in manifest.declared_losses:
            row = loss_table.add_row().cells
            row[0].text = str(loss.get("mappingSpec", ""))
            row[1].text = str(loss.get("path", ""))
            row[2].text = str(loss.get("kind", ""))
            row[3].text = str(loss.get("detail", ""))
    else:
        document.add_paragraph("None.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_xlsx(manifest: WorkshopPackManifest) -> bytes:
    workbook = Workbook()
    files_sheet = workbook.active
    assert files_sheet is not None
    files_sheet.title = "Files"
    files_sheet.append(["Path", "SHA-256", "Description"])
    for pack_file in manifest.files:
        files_sheet.append([pack_file.path, pack_file.sha256, pack_file.description])

    losses_sheet = workbook.create_sheet("Declared losses")
    losses_sheet.append(["Mapping spec", "Path", "Kind", "Detail"])
    for loss in manifest.declared_losses:
        losses_sheet.append(
            [loss.get("mappingSpec", ""), loss.get("path", ""), loss.get("kind", ""), loss.get("detail", "")]
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
